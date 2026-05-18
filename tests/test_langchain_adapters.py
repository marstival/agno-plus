"""Tests for LangChain adapter loaders.

langchain-core is an optional dep. These tests mock it so the suite runs
without installing it. The mock FakeDocument mimics langchain_core.documents.Document.
"""

from __future__ import annotations

import csv
import io
import sys
import tempfile
import os
from dataclasses import dataclass
from types import ModuleType
from typing import Any
from unittest.mock import patch

import pytest

from agno_plus.adapters.langchain.spreadsheet_loader import LangChainSpreadsheetLoader
from agno_plus.adapters.langchain.audio_loader import LangChainAudioLoader
from agno_plus.adapters.langchain.image_loader import LangChainImageLoader


# ---------------------------------------------------------------------------
# Fake langchain_core shim injected via sys.modules
# ---------------------------------------------------------------------------


@dataclass
class FakeDocument:
    page_content: str
    metadata: dict[str, Any]


def _inject_fake_langchain() -> None:
    """Insert a minimal langchain_core shim into sys.modules."""
    if "langchain_core" in sys.modules:
        return

    pkg = ModuleType("langchain_core")
    docs_mod = ModuleType("langchain_core.documents")
    docs_mod.Document = FakeDocument  # type: ignore[attr-defined]
    pkg.documents = docs_mod  # type: ignore[attr-defined]
    sys.modules["langchain_core"] = pkg
    sys.modules["langchain_core.documents"] = docs_mod


_inject_fake_langchain()


# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------


def make_csv_file(rows: list[list[str]]) -> str:
    """Write rows to a temp CSV file and return the path."""
    buf = io.StringIO()
    csv.writer(buf).writerows(rows)
    tmp = tempfile.NamedTemporaryFile(suffix=".csv", delete=False, mode="w")
    tmp.write(buf.getvalue())
    tmp.close()
    return tmp.name


def make_xlsx_file(rows: list[list[str]]) -> str:
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    for row in rows:
        ws.append(row)
    tmp = tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False)
    wb.save(tmp.name)
    tmp.close()
    return tmp.name


# ---------------------------------------------------------------------------
# LangChainSpreadsheetLoader
# ---------------------------------------------------------------------------


class TestSpreadsheetLoader:
    def setup_method(self):
        self.csv_path = make_csv_file([
            ["Name", "Score", "Grade"],
            ["Alice", "95", "A"],
            ["Bob", "82", "B"],
            ["Carol", "78", "C"],
            ["Dave", "91", "A"],
        ])

    def teardown_method(self):
        os.unlink(self.csv_path)

    def test_load_returns_list(self):
        loader = LangChainSpreadsheetLoader(self.csv_path)
        docs = loader.load()
        assert isinstance(docs, list)
        assert len(docs) >= 1

    def test_documents_are_fake_lc_type(self):
        loader = LangChainSpreadsheetLoader(self.csv_path)
        docs = loader.load()
        assert all(isinstance(d, FakeDocument) for d in docs)

    def test_page_content_has_table_data(self):
        loader = LangChainSpreadsheetLoader(self.csv_path)
        docs = loader.load()
        combined = " ".join(d.page_content for d in docs)
        assert "Alice" in combined
        assert "Bob" in combined

    def test_metadata_has_source(self):
        loader = LangChainSpreadsheetLoader(self.csv_path)
        docs = loader.load()
        for doc in docs:
            assert doc.metadata["source"] == self.csv_path
            assert doc.metadata["source_type"] == "spreadsheet"

    def test_metadata_has_block_type(self):
        loader = LangChainSpreadsheetLoader(self.csv_path)
        docs = loader.load()
        for doc in docs:
            assert "block_type" in doc.metadata

    def test_lazy_load_is_iterator(self):
        loader = LangChainSpreadsheetLoader(self.csv_path)
        it = loader.lazy_load()
        first = next(it)
        assert isinstance(first, FakeDocument)

    def test_xlsx_file_works(self):
        xlsx_path = make_xlsx_file([
            ["Product", "Price", "Units"],
            ["Widget", "9.99", "100"],
            ["Gadget", "19.99", "50"],
            ["Doohickey", "4.99", "200"],
        ])
        try:
            loader = LangChainSpreadsheetLoader(xlsx_path)
            docs = loader.load()
            assert len(docs) >= 1
            combined = " ".join(d.page_content for d in docs)
            assert "Widget" in combined
        finally:
            os.unlink(xlsx_path)

    def test_missing_langchain_raises_import_error(self):
        # Temporarily remove the shim to test the error path
        saved = sys.modules.pop("langchain_core", None)
        saved_docs = sys.modules.pop("langchain_core.documents", None)
        try:
            loader = LangChainSpreadsheetLoader(self.csv_path)
            with pytest.raises(ImportError, match="langchain-core"):
                loader.load()
        finally:
            if saved:
                sys.modules["langchain_core"] = saved
            if saved_docs:
                sys.modules["langchain_core.documents"] = saved_docs


# ---------------------------------------------------------------------------
# LangChainAudioLoader (mocks the core AudioReader to avoid Whisper dep)
# ---------------------------------------------------------------------------


class TestAudioLoader:
    def setup_method(self):
        self.audio_path = tempfile.NamedTemporaryFile(
            suffix=".mp3", delete=False
        )
        self.audio_path.write(b"fake audio bytes")
        self.audio_path.close()

    def teardown_method(self):
        os.unlink(self.audio_path.name)

    def _make_loader(self) -> LangChainAudioLoader:
        return LangChainAudioLoader(self.audio_path.name)

    def test_load_calls_core_reader(self):
        from agno_plus.core.models import Document

        fake_doc = Document(
            id="a1", content="transcript text", source_type="audio", source_name="test.mp3"
        )
        with patch(
            "agno_plus.adapters.langchain.audio_loader.AudioReader.read",
            return_value=[fake_doc],
        ):
            docs = self._make_loader().load()

        assert len(docs) == 1
        assert isinstance(docs[0], FakeDocument)
        assert docs[0].page_content == "transcript text"

    def test_metadata_source_type(self):
        from agno_plus.core.models import Document

        fake_doc = Document(
            id="a2", content="audio content", source_type="audio", source_name="clip.mp3"
        )
        with patch(
            "agno_plus.adapters.langchain.audio_loader.AudioReader.read",
            return_value=[fake_doc],
        ):
            docs = self._make_loader().load()

        assert docs[0].metadata["source_type"] == "audio"
        assert docs[0].metadata["source"] == self.audio_path.name

    def test_lazy_load_yields(self):
        from agno_plus.core.models import Document

        fake_doc = Document(id="a3", content="words", source_type="audio", source_name="x.mp3")
        with patch(
            "agno_plus.adapters.langchain.audio_loader.AudioReader.read",
            return_value=[fake_doc],
        ):
            it = self._make_loader().lazy_load()
            first = next(it)
        assert first.page_content == "words"


# ---------------------------------------------------------------------------
# LangChainImageLoader (mocks the core ImageReader to avoid OpenAI dep)
# ---------------------------------------------------------------------------


class TestImageLoader:
    def setup_method(self):
        self.image_path = tempfile.NamedTemporaryFile(suffix=".jpg", delete=False)
        self.image_path.write(b"fake image bytes")
        self.image_path.close()

    def teardown_method(self):
        os.unlink(self.image_path.name)

    def test_load_calls_core_reader(self):
        from agno_plus.core.models import Document

        fake_doc = Document(
            id="i1", content="OCR extracted text", source_type="image", source_name="receipt.jpg"
        )
        with patch(
            "agno_plus.adapters.langchain.image_loader.ImageReader.read",
            return_value=[fake_doc],
        ):
            loader = LangChainImageLoader(self.image_path.name)
            docs = loader.load()

        assert len(docs) == 1
        assert docs[0].page_content == "OCR extracted text"
        assert docs[0].metadata["source_type"] == "image"

    def test_metadata_source_path(self):
        from agno_plus.core.models import Document

        fake_doc = Document(id="i2", content="text", source_type="image", source_name="img.jpg")
        with patch(
            "agno_plus.adapters.langchain.image_loader.ImageReader.read",
            return_value=[fake_doc],
        ):
            loader = LangChainImageLoader(self.image_path.name)
            docs = loader.load()

        assert docs[0].metadata["source"] == self.image_path.name

    def test_lazy_load_yields(self):
        from agno_plus.core.models import Document

        fake_doc = Document(id="i3", content="receipt", source_type="image", source_name="r.jpg")
        with patch(
            "agno_plus.adapters.langchain.image_loader.ImageReader.read",
            return_value=[fake_doc],
        ):
            it = LangChainImageLoader(self.image_path.name).lazy_load()
            first = next(it)
        assert first.page_content == "receipt"
